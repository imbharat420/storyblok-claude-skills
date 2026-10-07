"""python suggestions_template.py rows.json SUGGESTIONS.xlsx  — rows.json: [{area,observation,recommendation,effort,impact,dependency,test,owner,notes}]"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
rows = json.load(open(sys.argv[1])); wb = Workbook(); wb.remove(wb.active)
thin = Side(style="thin", color="000000"); border = Border(left=thin,right=thin,top=thin,bottom=thin)
hdr = PatternFill("solid", fgColor="D9D9D9"); alt = PatternFill("solid", fgColor="F2F2F2")
H = ["ID","Area","Observation","Recommendation","Effort","Impact","Dependency","Falsification test","Decision","Owner","Notes"]
def sheet(name, items):
    ws = wb.create_sheet(name); ws.append(H)
    dv = DataValidation(type="list", formula1='"Accept,Defer,Reject"', allow_blank=True); ws.add_data_validation(dv)
    for i, r in enumerate(items, 1):
        ws.append([f"{name[:3].upper()}-{i:03d}", r["area"], r["observation"], r["recommendation"], r["effort"], r["impact"], r.get("dependency",""), r["test"], "", r.get("owner",""), r.get("notes","")])
        dv.add(ws.cell(row=i+1, column=9))
    for i,row in enumerate(ws.iter_rows(), start=1):
        for c in row:
            c.font = Font(name="Arial", size=10, bold=(i==1)); c.border = border
            if i == 1: c.fill = hdr
            elif i % 2 == 0: c.fill = alt
    for col,w in zip("ABCDEFGHIJK",[10,14,50,50,8,8,18,40,10,14,30]): ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"; ws.page_setup.orientation = "landscape"; ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToWidth = 1
order = lambda r: (-int(r["impact"]), "SMLX".index(r["effort"][0]))
quick = sorted([r for r in rows if int(r["impact"]) >= 4 and r["effort"][0] in "SM"], key=order)
sheet("Summary", quick)
for area in ["UI","Functionality","Storyblok","SEO","Performance","Accessibility"]:
    sheet(area, sorted([r for r in rows if r["area"] == area], key=order))
wb.save(sys.argv[2]); print("written", sys.argv[2], len(rows), "rows")
