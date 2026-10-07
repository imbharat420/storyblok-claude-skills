"""Build ESTIMATE.xlsx in house style from counts.json. Usage: python estimate_template.py counts.json out.xlsx"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
BANDS = {"token_set":(6,10,16),"primitive":(2,4,8),"block_static":(3,5,8),"block_interactive":(6,10,18),"page_template":(4,8,14),
         "schema_push":(0.5,1,2),"seed_content_type":(2,4,10),"language":(8,16,32),"integration":(8,20,48),"migration_entity":(4,10,24),"qa_page_viewport":(0.5,1,2)}
counts = json.load(open(sys.argv[1])); fixed = counts.get("fixed_hours", {}); conf = counts.get("confidence_pct", 15)
wb = Workbook(); ws = wb.active; ws.title = "Estimate"
thin = Side(style="thin", color="000000"); border = Border(left=thin,right=thin,top=thin,bottom=thin)
hdr = PatternFill("solid", fgColor="D9D9D9"); alt = PatternFill("solid", fgColor="F2F2F2"); font = Font(name="Arial", size=10)
rows = [["Item class","Count","Low (h)","Typical (h)","High (h)","Low total","Typical total","High total"]]
for k,(lo,ty,hi) in BANDS.items():
    n = counts.get(k, 0); rows.append([k, n, lo, ty, hi, n*lo, n*ty, n*hi])
for k,v in fixed.items(): rows.append([f"fixed: {k}", 1, v, v, v, v, v, v])
for r in rows: ws.append(r)
n = len(rows)+1
ws.append(["TOTAL","","","","",f"=SUM(F2:F{n-1})",f"=SUM(G2:G{n-1})",f"=SUM(H2:H{n-1})"])
ws.append(["Confidence band", f"+/-{conf}%","","","","","",""])
for i,row in enumerate(ws.iter_rows(), start=1):
    for c in row:
        c.font = Font(name="Arial", size=10, bold=(i==1)); c.border = border
        if i == 1: c.fill = hdr
        elif i % 2 == 0: c.fill = alt
for col,w in zip("ABCDEFGH",[28,8,10,12,10,12,14,12]): ws.column_dimensions[col].width = w
ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(sys.argv[2]); print("written", sys.argv[2])
